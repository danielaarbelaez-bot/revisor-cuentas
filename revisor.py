import io
import re
import zipfile
import unicodedata
import datetime as dt
from collections import defaultdict

import pandas as pd
import streamlit as st
import holidays
from docx import Document
from pypdf import PdfReader
from openpyxl import load_workbook

st.title('Revisor de cuentas de cobro | Demostración V4')
st.warning('Solo para pruebas con datos ficticios. No cargue documentos institucionales ni datos personales reales.')

MAX_FILE = 20 * 1024 * 1024
MAX_TOTAL = 80 * 1024 * 1024
MAX_ZIP_FILES = 40


def normalizar(valor):
    texto = unicodedata.normalize('NFKD', str(valor if valor is not None else ''))
    return ''.join(c for c in texto if not unicodedata.combining(c)).upper().strip()


@st.cache_data
def base_demo():
    return pd.DataFrame([
        {'NOMBRE': 'ANA MARIA PRUEBA', 'CEDULA': '999000001', 'CONTRATO': 'DEF-DEMO-001', 'REGIONAL': 'ANTIOQUIA', 'PROYECTO': 'EXIGIBILIDAD', 'HONORARIOS': 5900000},
        {'NOMBRE': 'LAURA PATRICIA EJEMPLO', 'CEDULA': '999000002', 'CONTRATO': 'DEF-DEMO-002', 'REGIONAL': 'QUINDIO', 'PROYECTO': 'VICTIMAS', 'HONORARIOS': 5900000},
    ])


def documentos(archivos):
    salida, total = [], 0
    for archivo in archivos:
        contenido = archivo.getvalue()
        if len(contenido) > MAX_FILE:
            st.error(f'{archivo.name}: supera el límite de 20 MB.')
            continue
        if archivo.name.lower().endswith('.zip'):
            try:
                with zipfile.ZipFile(io.BytesIO(contenido)) as z:
                    candidatos = [x for x in z.infolist() if not x.is_dir()]
                    if len(candidatos) > MAX_ZIP_FILES:
                        st.error(f'{archivo.name}: contiene demasiados archivos.')
                        continue
                    for info in candidatos:
                        nombre = info.filename.replace('\\', '/')
                        if ('..' in nombre.split('/') or nombre.startswith('/') or info.flag_bits & 0x1 or
                            not nombre.lower().endswith(('.pdf', '.docx', '.xlsx', '.xlsm'))):
                            continue
                        if info.file_size > MAX_FILE or total + info.file_size > MAX_TOTAL:
                            st.warning('Se omitió un archivo por límites de tamaño.')
                            continue
                        data = z.read(info)
                        total += len(data)
                        salida.append((nombre, data))
            except (zipfile.BadZipFile, RuntimeError, ValueError, OSError):
                st.error(f'No se pudo abrir el ZIP {archivo.name}.')
        elif archivo.name.lower().endswith(('.pdf', '.docx', '.xlsx', '.xlsm')):
            if total + len(contenido) <= MAX_TOTAL:
                salida.append((archivo.name, contenido))
                total += len(contenido)
    return salida


def extraer(nombre, contenido):
    ext = nombre.lower().rsplit('.', 1)[-1]
    datos = {'texto': '', 'ubicaciones': [], 'hojas': {}, 'version_f68': None, 'error': None, 'escaneo': False}
    try:
        if ext == 'pdf':
            pdf = PdfReader(io.BytesIO(contenido))
            for i, pagina in enumerate(pdf.pages):
                texto = pagina.extract_text() or ''
                datos['ubicaciones'].append((f'página {i+1}', texto))
                if len(texto.strip()) < 20:
                    datos['escaneo'] = True
        elif ext == 'docx':
            doc = Document(io.BytesIO(contenido))
            for i, p in enumerate(doc.paragraphs):
                if p.text.strip():
                    datos['ubicaciones'].append((f'párrafo {i+1}', p.text))
            for t, tabla in enumerate(doc.tables):
                for i, fila in enumerate(tabla.rows):
                    for j, celda in enumerate(fila.cells):
                        if celda.text.strip():
                            datos['ubicaciones'].append((f'tabla {t+1}, fila {i+1}, columna {j+1}', celda.text))
            # Word may contain text in headers, footers and embedded images not extracted here.
        elif ext in ('xlsx', 'xlsm'):
            libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=False)
            for hoja in libro.worksheets:
                valores = {}
                for fila in hoja.iter_rows():
                    for celda in fila:
                        if celda.value is not None and str(celda.value).strip():
                            val = str(celda.value)
                            valores[celda.coordinate] = val
                            datos['ubicaciones'].append((f'hoja {hoja.title}, celda {celda.coordinate}', val))
                datos['hojas'][hoja.title] = valores
            if 'F68' in normalizar(nombre):
                for titulo, celdas in datos['hojas'].items():
                    if 'FORMULARIO PARA PAGOS' in normalizar(titulo):
                        datos['version_f68'] = celdas.get('AA4')
                        break
        datos['texto'] = '\n'.join(texto for _, texto in datos['ubicaciones'])
    except Exception as exc:
        datos['error'] = type(exc).__name__
    return datos


def identificar(analizados, base):
    texto = normalizar('\n'.join(x['texto'] for x in analizados.values()))
    candidatos = []
    for _, fila in base.iterrows():
        señales = []
        if normalizar(fila['CONTRATO']) in texto:
            señales.append('contrato')
        if re.search(r'(?<!\d)' + re.escape(str(fila['CEDULA'])) + r'(?!\d)', texto):
            señales.append('cédula')
        if normalizar(fila['NOMBRE']) in texto:
            señales.append('nombre')
        candidatos.append((len(señales), fila, señales))
    candidatos.sort(key=lambda v: v[0], reverse=True)
    if not candidatos or candidatos[0][0] < 2 or (len(candidatos) > 1 and candidatos[0][0] == candidatos[1][0]):
        return None, 'Identificación no concluyente: se requieren dos coincidencias y un resultado único.'
    return candidatos[0][1], 'Coincidencia preliminar mediante ' + ', '.join(candidatos[0][2]) + '.'


def es_marcador(valor):
    v = normalizar(valor).strip(' .:-_')
    return not v or v in {'N/A', 'NA', 'NO APLICA', 'ANONIMIZADO', 'RESERVADO', 'CONFIDENCIAL', 'XXXXX', 'XXXX', 'XXX', 'NOMBRE', 'APELLIDO', 'CEDULA', 'DOCUMENTO', 'TELEFONO', 'CORREO', 'DIRECCION', 'SIN DATO'} or bool(re.fullmatch(r'[X*#\-\s]{3,}', v))


def riesgos_privacidad(nombre, datos, es_matriz):
    avisos = []
    campos = re.compile(r'\b(NOMBRE(?:S)?|APELLIDO(?:S)?|C[EÉ]DULA|DOCUMENTO(?: DE IDENTIDAD)?|TEL[EÉ]FONO|CELULAR|CORREO(?: ELECTR[OÓ]NICO)?|DIRECCI[OÓ]N)\b', re.I)
    email = re.compile(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b')
    telefono = re.compile(r'(?<!\d)3\d{9}(?!\d)')
    identificacion = re.compile(r'(?<!\d)\d{7,10}(?!\d)')
    if es_matriz and datos['hojas']:
        for hoja, celdas in datos['hojas'].items():
            encabezados = [(pos, val) for pos, val in celdas.items() if campos.search(normalizar(val))]
            for pos, encabezado in encabezados:
                col = re.match(r'[A-Z]+', pos).group()
                fila_enc = int(re.search(r'\d+', pos).group())
                ejemplos = []
                for celda, valor in celdas.items():
                    if re.match(r'[A-Z]+', celda).group() != col or int(re.search(r'\d+', celda).group()) <= fila_enc:
                        continue
                    if not es_marcador(valor) and not str(valor).startswith('='):
                        ejemplos.append(celda)
                if ejemplos:
                    avisos.append((f'hoja {hoja}, columna {col}', f'La columna «{encabezado[:45]}» contiene posibles identificadores en {len(ejemplos)} celda(s). Revisar anonimización. No se muestran valores.'))
                else:
                    avisos.append((f'hoja {hoja}, columna {col}', f'Existe la columna «{encabezado[:45]}»; no se detectaron valores identificables en celdas posteriores. Comprobar necesidad de conservarla.'))
    # En soportes se alerta de patrones, sin mostrar los datos; las cédulas aisladas son ambiguas.
    if not ('F62' in normalizar(nombre) or 'F68' in normalizar(nombre) or 'PLANILLA' in normalizar(nombre)):
        for ubicacion, texto in datos['ubicaciones']:
            if email.search(texto) or telefono.search(texto):
                avisos.append((ubicacion, 'Posible correo o teléfono identificable. Verificar si corresponde a un tercero y si debe anonimizarse.'))
            elif not es_matriz and identificacion.search(texto) and campos.search(texto):
                avisos.append((ubicacion, 'Posible número de identificación asociado a un campo personal. Revisar anonimización.'))
    # Prevent repetitive alerts from large files.
    return avisos[:35]



# V4: cotejo conservador de campos identificables. No inventa valores faltantes.
MESES = {normalizar(m): i for i, m in enumerate(['Enero','Febrero','Marzo','Abril','Mayo','Junio','Julio','Agosto','Septiembre','Octubre','Noviembre','Diciembre'], 1)}


def fecha_partes(dia, mes, anio):
    try:
        m = MESES.get(normalizar(mes), None)
        if m is None and str(mes).strip().isdigit():
            m = int(mes)
        return dt.date(int(str(anio).strip()), m, int(str(dia).strip()))
    except (ValueError, TypeError):
        return None


def valor_celda(datos, coord, hoja='Formulario para pagos'):
    for nombre, celdas in datos['hojas'].items():
        if normalizar(nombre) == normalizar(hoja):
            return celdas.get(coord)
    return None


def limpiar_campo(valor):
    v = normalizar(valor)
    return re.sub(r'\s+', ' ', v).strip()


def f68_campos(datos):
    if not any(normalizar(h) == 'FORMULARIO PARA PAGOS' for h in datos['hojas']):
        return {}
    v = lambda c: valor_celda(datos, c)
    return {
        'contrato': v('K17'), 'nombre': v('L19'), 'cedula': v('AB19'),
        'inicio': fecha_partes(v('N13'),v('P13'),v('T13')),
        'fin': fecha_partes(v('X13'),v('Z13'),v('AD13')),
        'mes_seguridad': v('AC102'),
        'salud': v('Z106'), 'pension': v('Z108'), 'arl': v('Z110'),
        'supervisor': v('J49'), 'valor_total': v('J41'),
    }


def f62_campos_bytes(contenido):
    try:
        doc = Document(io.BytesIO(contenido))
        # Estructura basada en la tabla de información básica del formato F62 V2.
        tabla = next((t for t in doc.tables if len(t.rows) > 13 and
                       'INFORMACION BASICA' in normalizar(t.rows[0].cells[0].text)), None)
        if tabla is None:
            return {}
        def dato(fila):
            if len(tabla.rows) <= fila:
                return None
            celdas = tabla.rows[fila].cells
            for c in celdas[4:]:
                if c.text.strip():
                    return c.text.strip()
            return None
        return {'contrato': dato(4), 'nombre': dato(5), 'cedula': dato(6),
                'supervisor': dato(13), 'valor_total': dato(10),
                'periodo_texto': dato(2)}
    except Exception:
        return {}


def cotejar_campos(nombre62, datos62, bytes62, nombre68, datos68, contratista, periodo, agregar):
    a, b = f62_campos_bytes(bytes62), f68_campos(datos68)
    if not a or not b:
        agregar('Comparación F62–F68', 'Revisión manual',
                'No fue posible leer ambos formatos con la estructura oficial. El cotejo de campos queda pendiente.')
        return
    for campo, titulo in [('contrato','número de contrato'), ('nombre','nombre de la contratista'),
                          ('cedula','documento de identidad')]:
        va, vb = a.get(campo), b.get(campo)
        if va and vb and not es_marcador(va) and not es_marcador(vb):
            if campo == 'cedula':
                coinciden = re.sub(r'\D','',str(va)) == re.sub(r'\D','',str(vb))
            else:
                coinciden = limpiar_campo(va) == limpiar_campo(vb)
            if not coinciden:
                agregar('Comparación F62–F68', 'Corregir',
                        f'El {titulo} no coincide entre F62 y F68. Ajustar el dato incorrecto para que ambos formatos sean consistentes, previa verificación del contrato.',
                        'F62: información básica / F68: formulario para pagos')
        else:
            agregar('Comparación F62–F68', 'Revisión manual',
                    f'No se pudo cotejar el {titulo} en ambos formatos; verificar los campos correspondientes.')
    inicio, fin = b.get('inicio'), b.get('fin')
    esperado_inicio = dt.date(periodo.year,periodo.month,1)
    esperado_fin = dt.date(periodo.year + (periodo.month == 12),periodo.month % 12 + 1,1) - dt.timedelta(days=1)
    if inicio and fin:
        if inicio != esperado_inicio or fin != esperado_fin:
            agregar(nombre68,'Corregir',
                    f'El período del F68 ({inicio:%d/%m/%Y} a {fin:%d/%m/%Y}) difiere del mes seleccionado ({esperado_inicio:%d/%m/%Y} a {esperado_fin:%d/%m/%Y}). Verificar el período aplicable y ajustar.',
                    'Formulario para pagos, fila 13')
    else:
        agregar(nombre68,'Revisión manual','No se pudo interpretar el período completo del F68 en la fila 13.')
    periodo62 = normalizar(a.get('periodo_texto') or '')
    if periodo62 and 'DEL' in periodo62 and not re.search(r'\b\d{1,2}\b',periodo62):
        agregar(nombre62,'Revisión manual','El período del F62 no tiene fechas legibles para cotejar.', 'Información básica, período de reporte')
    mes_ss = b.get('mes_seguridad')
    if mes_ss and normalizar(mes_ss) in MESES and MESES[normalizar(mes_ss)] != periodo.month:
        agregar(nombre68,'Corregir',
                'El mes de aportes de seguridad social registrado en F68 no corresponde al mes seleccionado. Verificar el período cotizado y corregir el campo si procede.',
                'Formulario para pagos!AC102')
    elif not mes_ss:
        agregar(nombre68,'Revisión manual','No se pudo comprobar el mes de aportes de seguridad social en F68.', 'Formulario para pagos!AC102')
    # No comparar honorarios o aportes cuando el documento no permite un cotejo inequívoco.
    agregar('Comparación F62–F68','Revisión manual',
            'Revisar valores de honorarios, aportes, obligaciones y soportes con documentos originales; el sistema no dispone de una verificación concluyente.')


st.success('Base ficticia incorporada automáticamente. No necesitas subir un Excel.')
with st.expander('Ver contratistas ficticias para la prueba'):
    st.dataframe(base_demo(), hide_index=True, use_container_width=True)

periodo = st.date_input('Mes de la cuenta', value=dt.date(2026, 10, 1))
archivos = st.file_uploader('Cargar F62, F68, planilla, matrices o ZIP de demostración', type=['pdf', 'docx', 'xlsx', 'xlsm', 'zip'], accept_multiple_files=True)

if st.button('Revisar cuenta automáticamente', type='primary', use_container_width=True):
    if not archivos:
        st.warning('Primero carga documentos de demostración.')
        st.stop()
    contenidos = documentos(archivos)
    if not contenidos:
        st.error('No se encontraron archivos compatibles.')
        st.stop()
    analizados = {nombre: extraer(nombre, contenido) for nombre, contenido in contenidos}
    contratista, mensaje = identificar(analizados, base_demo())
    if contratista is None:
        st.warning(mensaje)
    else:
        st.success('Contratista identificada: ' + str(contratista['NOMBRE']))
        st.caption(f"Contrato: {contratista['CONTRATO']} | Regional: {contratista['REGIONAL']} | Proyecto: {contratista['PROYECTO']}")
        st.caption(mensaje)

    hallazgos = []
    def agregar(documento, estado, observacion, ubicacion='—'):
        hallazgos.append({'Documento': documento, 'Estado': estado, 'Ubicación': ubicacion, 'Observación': observacion})

    nombres = [normalizar(n) for n, _ in contenidos]
    requeridos = {
        'F62': any('F62' in n for n in nombres),
        'F68': any('F68' in n for n in nombres),
        'Planilla PDF': any('PLANILLA' in n and n.endswith('.PDF') for n in nombres),
        'Matriz de casos': any('CASO' in n and n.endswith(('.XLSX', '.XLSM')) for n in nombres),
        'Matriz PYD Excel': any('PYD' in n and n.endswith(('.XLSX', '.XLSM')) for n in nombres),
    }
    for etiqueta, existe in requeridos.items():
        if not existe:
            agregar('Expediente', 'Verificar', f'No se identificó {etiqueta}; revisar si el archivo tiene otro nombre.')

    f62, f68 = [], []
    for nombre, datos in analizados.items():
        clave = normalizar(nombre)
        plano = normalizar(datos['texto'])
        if datos['error']:
            agregar(nombre, 'Revisión manual', 'No fue posible leer el documento completamente; verificar archivo o protección.')
            continue
        if not datos['texto'].strip():
            agregar(nombre, 'Revisión manual', 'No se extrajo texto; podría ser un escaneo o documento protegido.')
        if datos['escaneo']:
            agregar(nombre, 'Revisión manual', 'Algunas páginas carecen de texto extraíble. Las imágenes y firmas no fueron verificadas.')
        if 'F62' in clave:
            f62.append((nombre, datos))
        if 'F68' in clave:
            f68.append((nombre, datos))
            if datos['version_f68'] is None:
                agregar(nombre, 'Revisión manual', 'No se encontró la celda de versión AA4 en la hoja «Formulario para pagos»; comprobar plantilla.')
            elif not re.search(r'(?<!\d)0?3(?:\.0)?(?!\d)', str(datos['version_f68'])):
                agregar(nombre, 'Corregir', f'La versión indicada en AA4 no corresponde a la versión 3 esperada.', 'Formulario para pagos!AA4')
        if 'F62' in clave or 'F68' in clave:
            agregar(nombre, 'Revisión manual', 'Comprobar firmas y correspondencia entre obligaciones, actividades y soportes; no se certifican automáticamente.')
            if contratista is not None and normalizar(contratista['CONTRATO']) not in plano:
                agregar(nombre, 'Verificar', 'No se encontró el contrato esperado en el texto extraído; comprobar el campo del formato.')
        if 'PLANILLA' in clave and clave.endswith('.PDF'):
            agregar(nombre, 'Revisión manual', 'Validar IBC, novedades, aportes, período cotizado y fecha de pago contra la planilla original. El valor de $2.360.000 no es universal.')
        es_matriz = clave.endswith(('.XLSX', '.XLSM')) and ('CASO' in clave or 'PYD' in clave)
        if es_matriz or ('SOPORTE' in clave or 'EVIDENCIA' in clave or 'ANEXO' in clave):
            for ubicacion, aviso in riesgos_privacidad(nombre, datos, es_matriz):
                agregar(nombre, 'Privacidad: verificar', aviso, ubicacion)

    if len(f62) == 1 and len(f68) == 1:
        nombre62, datos62 = f62[0]
        nombre68, datos68 = f68[0]
        bytes62 = next((b for n,b in contenidos if n == nombre62), None)
        if nombre62.lower().endswith('.docx') and bytes62:
            cotejar_campos(nombre62, datos62, bytes62, nombre68, datos68, contratista, periodo, agregar)
        else:
            agregar('Comparación F62–F68','Revisión manual',
                    'El cotejo estructurado de F62 requiere un archivo Word DOCX con la plantilla oficial.')
    elif len(f62) > 1 or len(f68) > 1:
        agregar('Comparación F62–F68', 'Verificar', 'Se recibieron varios F62 o F68; seleccionar el par correspondiente antes de comparar.')

    siguiente = dt.date(periodo.year + (periodo.month == 12), periodo.month % 12 + 1, 1)
    ultimo = siguiente - dt.timedelta(days=1)
    festivos = holidays.country_holidays('CO', years=[siguiente.year, siguiente.year + 1])
    fecha = siguiente
    while fecha.weekday() >= 5 or fecha in festivos:
        fecha += dt.timedelta(days=1)
    st.info(f'Período seleccionado: 01/{periodo.month:02d}/{periodo.year} a {ultimo:%d/%m/%Y}. Primer día hábil siguiente: {fecha:%d/%m/%Y} (referencia; confirmar plazo contractual).')

    if not hallazgos:
        agregar('Expediente', 'Revisión manual', 'Sin alertas automáticas; no equivale a validación integral.')
    tabla = pd.DataFrame(hallazgos).drop_duplicates()
    st.subheader('Observaciones preliminares')
    st.dataframe(tabla, hide_index=True, use_container_width=True)
    st.caption('Las alertas de privacidad son indicios, no pruebas concluyentes. No se verifican imágenes, escaneos, firmas ni datos ocultos de Excel.')

    # Solo los hallazgos confirmados y corregibles se incluyen en el correo.
    # Alertas de privacidad, dudas y validaciones manuales quedan para revisión interna.
    correcciones = tabla[tabla['Estado'].eq('Corregir')].copy()
    pendientes = tabla[~tabla['Estado'].eq('Corregir')].copy()

    st.subheader('Correcciones comprobadas para enviar a la contratista')
    if correcciones.empty:
        st.success('No se identificaron correcciones comprobadas en los campos que esta versión puede analizar. Esto no equivale a aprobar la cuenta.')
    else:
        st.dataframe(correcciones, hide_index=True, use_container_width=True)

    st.subheader('Verificaciones internas (no incluir en el correo)')
    if pendientes.empty:
        st.info('No se generaron verificaciones internas adicionales.')
    else:
        st.dataframe(pendientes, hide_index=True, use_container_width=True)

    correo = 'Cordial saludo,\n\n'
    if correcciones.empty:
        correo += ('En la revisión automática preliminar no se identificaron correcciones comprobadas '
                   'para solicitar. La cuenta permanece pendiente de la revisión manual correspondiente.\n\n')
    else:
        correo += ('De acuerdo con la revisión realizada a la cuenta de cobro, '
                   'solicito efectuar las siguientes correcciones identificadas en los documentos: '
                   '\n\n')
        for documento, grupo in correcciones.groupby('Documento', sort=False):
            correo += documento + '\n'
            for _, fila in grupo.iterrows():
                ubic = '' if fila['Ubicación'] == '—' else f" ({fila['Ubicación']})"
                correo += f"- {fila['Observación']}{ubic}\n"
            correo += '\n'
        correo += 'Agradezco remitir la documentación ajustada para continuar con el trámite.\n\n'
    correo += 'Cordialmente,'
    st.text_area('Correo único para Outlook — solo correcciones comprobadas', correo, height=280)
    st.download_button('Descargar correo para Outlook', correo, 'correo_correcciones_demo_v4.txt', mime='text/plain')
    st.caption('No se envían automáticamente correos. Las alertas de privacidad se excluyen del correo hasta confirmar que corresponden a datos personales que deben anonimizarse.')
    st.warning('Demostración V4: coteja campos estructurados cuando se usan las plantillas oficiales. No certifica anonimización completa, firmas, macros, pagos ni obligaciones. No apta para expedientes reales.')
